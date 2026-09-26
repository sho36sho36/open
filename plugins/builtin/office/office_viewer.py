from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtGui import QFont
from PySide6.QtWidgets import (
    QAbstractItemView,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)


class OfficeViewer(QWidget):
    """Office / OpenDocumentファイル表示用ビューア。"""

    SPREADSHEET_EXTENSIONS = {
        ".xlsx",
        ".ods",
    }

    DOCUMENT_EXTENSIONS = {
        ".docx",
        ".odt",
        ".rtf",
    }

    PRESENTATION_EXTENSIONS = {
        ".pptx",
        ".odp",
    }

    def __init__(
        self,
        file_info,
        parent=None,
    ):
        super().__init__(parent)

        self.file_info = file_info
        self.file_path = Path(file_info.path)

        self.setWindowTitle(
            f"{file_info.name} - open"
        )

        self._build_ui()
        self._open_file()

    # =========================================================
    # UI
    # =========================================================

    def _build_ui(self):
        self.title_label = QLabel(
            self.file_info.name
        )

        font = QFont()
        font.setBold(True)
        font.setPointSize(12)

        self.title_label.setFont(font)

        # ここが以前のTypeErrorの修正版
        self.type_label = QLabel()

        self.type_label.setAlignment(
            Qt.AlignmentFlag.AlignRight
            | Qt.AlignmentFlag.AlignVCenter
        )

        self.reload_button = QPushButton(
            "🔄 再読み込み"
        )

        self.reload_button.clicked.connect(
            self._reload
        )

        header = QHBoxLayout()

        header.addWidget(
            self.title_label
        )

        header.addStretch(1)

        header.addWidget(
            self.type_label
        )

        header.addWidget(
            self.reload_button
        )

        self.content = QWidget()

        self.content_layout = QVBoxLayout(
            self.content
        )

        self.content_layout.setContentsMargins(
            0,
            0,
            0,
            0,
        )

        layout = QVBoxLayout(self)

        layout.addLayout(header)

        layout.addWidget(
            self.content,
            1,
        )

    def _clear_content(self):
        while self.content_layout.count():
            item = self.content_layout.takeAt(0)

            widget = item.widget()

            if widget is not None:
                widget.deleteLater()

    # =========================================================
    # Open
    # =========================================================

    def _open_file(self):
        suffix = self.file_path.suffix.lower()

        if suffix in self.SPREADSHEET_EXTENSIONS:
            self._open_spreadsheet()
            return

        if suffix in self.DOCUMENT_EXTENSIONS:
            self._open_document()
            return

        if suffix in self.PRESENTATION_EXTENSIONS:
            self._open_presentation()
            return

        self._show_error(
            "対応していないOffice形式です。"
        )

    # =========================================================
    # Spreadsheet
    # =========================================================

    def _open_spreadsheet(self):
        self.type_label.setText(
            "📊 Spreadsheet"
        )

        self._clear_content()

        suffix = self.file_path.suffix.lower()

        if suffix == ".xlsx":
            self._open_xlsx()
            return

        if suffix == ".ods":
            self._open_ods()
            return

    def _open_xlsx(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self._show_error(
                "openpyxlがインストールされていません。\n\n"
                "pip install openpyxl"
            )
            return

        try:
            workbook = load_workbook(
                self.file_path,
                read_only=True,
                data_only=False,
            )

            tabs = QTabWidget()

            for sheet in workbook.worksheets:
                table = self._create_table()

                row_data = []

                for row in sheet.iter_rows(
                    values_only=True
                ):
                    row_data.append(
                        [
                            ""
                            if value is None
                            else str(value)
                            for value in row
                        ]
                    )

                self._fill_table(
                    table,
                    row_data,
                )

                tabs.addTab(
                    table,
                    sheet.title,
                )

            self.content_layout.addWidget(
                tabs
            )

            workbook.close()

        except Exception as exc:
            self._show_error(
                "Excelファイルを開けませんでした。\n\n"
                f"{exc}"
            )

    def _open_ods(self):
        try:
            from odf import teletype
            from odf.opendocument import load
            from odf.table import (
                Table,
                TableCell,
                TableRow,
            )
        except ImportError:
            self._show_error(
                "odfpyがインストールされていません。\n\n"
                "pip install odfpy"
            )
            return

        try:
            document = load(
                str(self.file_path)
            )

            tables = (
                document.spreadsheet.getElementsByType(
                    Table
                )
            )

            tabs = QTabWidget()

            for sheet in tables:
                table_widget = self._create_table()

                data = []

                rows = sheet.getElementsByType(
                    TableRow
                )

                for row in rows:
                    values = []

                    cells = row.getElementsByType(
                        TableCell
                    )

                    for cell in cells:
                        values.append(
                            teletype.extractText(
                                cell
                            )
                        )

                    data.append(values)

                self._fill_table(
                    table_widget,
                    data,
                )

                sheet_name = (
                    sheet.getAttribute("name")
                    or "Sheet"
                )

                tabs.addTab(
                    table_widget,
                    sheet_name,
                )

            self.content_layout.addWidget(
                tabs
            )

        except Exception as exc:
            self._show_error(
                "ODSファイルを開けませんでした。\n\n"
                f"{exc}"
            )

    # =========================================================
    # Document
    # =========================================================

    def _open_document(self):
        self.type_label.setText(
            "📝 Document"
        )

        self._clear_content()

        suffix = self.file_path.suffix.lower()

        if suffix == ".docx":
            self._open_docx()
            return

        if suffix == ".odt":
            self._open_odt()
            return

        if suffix == ".rtf":
            self._open_rtf()
            return

    def _open_docx(self):
        try:
            from docx import Document
        except ImportError:
            self._show_error(
                "python-docxがインストールされていません。\n\n"
                "pip install python-docx"
            )
            return

        try:
            document = Document(
                self.file_path
            )

            editor = self._create_text_editor()

            text_parts = []

            for paragraph in document.paragraphs:
                text_parts.append(
                    paragraph.text
                )

            editor.setPlainText(
                "\n\n".join(text_parts)
            )

            self.content_layout.addWidget(
                editor
            )

        except Exception as exc:
            self._show_error(
                "Wordファイルを開けませんでした。\n\n"
                f"{exc}"
            )

    def _open_odt(self):
        try:
            from odf import teletype
            from odf.opendocument import load
            from odf.text import P
        except ImportError:
            self._show_error(
                "odfpyがインストールされていません。\n\n"
                "pip install odfpy"
            )
            return

        try:
            document = load(
                str(self.file_path)
            )

            editor = self._create_text_editor()

            text_parts = []

            paragraphs = (
                document.getElementsByType(P)
            )

            for paragraph in paragraphs:
                text_parts.append(
                    teletype.extractText(
                        paragraph
                    )
                )

            editor.setPlainText(
                "\n\n".join(text_parts)
            )

            self.content_layout.addWidget(
                editor
            )

        except Exception as exc:
            self._show_error(
                "ODTファイルを開けませんでした。\n\n"
                f"{exc}"
            )

    def _open_rtf(self):
        try:
            text = self.file_path.read_text(
                encoding="utf-8",
                errors="replace",
            )

        except Exception as exc:
            self._show_error(
                "RTFファイルを読み込めませんでした。\n\n"
                f"{exc}"
            )
            return

        editor = self._create_text_editor()

        editor.setPlainText(text)

        self.content_layout.addWidget(
            editor
        )

    # =========================================================
    # Presentation
    # =========================================================

    def _open_presentation(self):
        self.type_label.setText(
            "📽️ Presentation"
        )

        self._clear_content()

        suffix = self.file_path.suffix.lower()

        if suffix == ".pptx":
            self._open_pptx()
            return

        if suffix == ".odp":
            self._open_odp()
            return

    def _open_pptx(self):
        try:
            from pptx import Presentation
        except ImportError:
            self._show_error(
                "python-pptxがインストールされていません。\n\n"
                "pip install python-pptx"
            )
            return

        try:
            presentation = Presentation(
                self.file_path
            )

            tabs = QTabWidget()

            for index, slide in enumerate(
                presentation.slides,
                start=1,
            ):
                editor = self._create_text_editor()

                lines = []

                for shape in slide.shapes:
                    if not hasattr(
                        shape,
                        "text",
                    ):
                        continue

                    text = shape.text.strip()

                    if text:
                        lines.append(text)

                if lines:
                    editor.setPlainText(
                        "\n\n".join(lines)
                    )
                else:
                    editor.setPlainText(
                        "(このスライドにはテキストがありません)"
                    )

                tabs.addTab(
                    editor,
                    f"Slide {index}",
                )

            self.content_layout.addWidget(
                tabs
            )

        except Exception as exc:
            self._show_error(
                "PowerPointファイルを開けませんでした。\n\n"
                f"{exc}"
            )

    def _open_odp(self):
        """
        ODPはodfpyの汎用XML要素から
        スライド単位でテキストを抽出します。

        odf.presentation.Pageは使用しません。
        """

        try:
            from odf import teletype
            from odf.opendocument import load
        except ImportError:
            self._show_error(
                "odfpyがインストールされていません。\n\n"
                "pip install odfpy"
            )
            return

        try:
            document = load(
                str(self.file_path)
            )

            tabs = QTabWidget()

            pages = []

            for element in document.presentation.childNodes:
                if (
                    getattr(
                        element,
                        "qname",
                        None,
                    )
                    == (
                        "urn:oasis:names:tc:opendocument:"
                        "xmlns:drawing:1.0",
                        "page",
                    )
                ):
                    pages.append(element)

            if not pages:
                # 名前空間に依存しすぎないフォールバック
                for element in document.presentation.childNodes:
                    if getattr(
                        element,
                        "tagName",
                        "",
                    ) == "draw:page":
                        pages.append(element)

            for index, page in enumerate(
                pages,
                start=1,
            ):
                text_parts = []

                self._collect_odp_text(
                    page,
                    text_parts,
                    teletype,
                )

                editor = self._create_text_editor()

                if text_parts:
                    editor.setPlainText(
                        "\n\n".join(
                            text_parts
                        )
                    )
                else:
                    editor.setPlainText(
                        "(このスライドにはテキストがありません)"
                    )

                tabs.addTab(
                    editor,
                    f"Slide {index}",
                )

            if tabs.count() == 0:
                editor = self._create_text_editor()

                editor.setPlainText(
                    "ODPのスライドを検出できませんでした。"
                )

                tabs.addTab(
                    editor,
                    "Presentation",
                )

            self.content_layout.addWidget(
                tabs
            )

        except Exception as exc:
            self._show_error(
                "ODPファイルを開けませんでした。\n\n"
                f"{exc}"
            )

    def _collect_odp_text(
        self,
        element,
        output,
        teletype,
    ):
        try:
            text = teletype.extractText(
                element
            )

            text = text.strip()

            if text and text not in output:
                output.append(text)

        except Exception:
            pass

        for child in getattr(
            element,
            "childNodes",
            [],
        ):
            self._collect_odp_text(
                child,
                output,
                teletype,
            )

    # =========================================================
    # Utility
    # =========================================================

    def _create_table(self):
        table = QTableWidget()

        table.setAlternatingRowColors(
            True
        )

        table.setEditTriggers(
            QAbstractItemView.EditTrigger.NoEditTriggers
        )

        table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectItems
        )

        table.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )

        table.horizontalHeader().setStretchLastSection(
            True
        )

        table.verticalHeader().setVisible(
            True
        )

        return table

    def _fill_table(
        self,
        table,
        rows,
    ):
        if not rows:
            return

        row_count = len(rows)

        column_count = max(
            len(row)
            for row in rows
        )

        if column_count == 0:
            return

        table.setRowCount(
            row_count
        )

        table.setColumnCount(
            column_count
        )

        for row_index, row in enumerate(
            rows
        ):
            for column_index, value in enumerate(
                row
            ):
                table.setItem(
                    row_index,
                    column_index,
                    QTableWidgetItem(
                        str(value)
                    ),
                )

        table.resizeColumnsToContents()

    def _create_text_editor(self):
        editor = QTextEdit()

        editor.setReadOnly(
            True
        )

        editor.setLineWrapMode(
            QTextEdit.LineWrapMode.WidgetWidth
        )

        return editor

    def _show_error(
        self,
        message,
    ):
        self._clear_content()

        label = QLabel(
            message
        )

        label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        label.setWordWrap(
            True
        )

        self.content_layout.addWidget(
            label
        )

    def _reload(self):
        self._clear_content()
        self._open_file()

    def closeEvent(
        self,
        event,
    ):
        event.accept()